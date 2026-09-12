#!/usr/bin/env python3
"""İnsan-robot işbirlikli vidalama senaryosunun SİSTEM TANIMI raporunu üretir.

    python3 docs/hrc_figurler.py          # figürler (bir kez yeterli)
    python3 docs/hrc_sistem_tanimi_uret.py

Bu belge bir gelişme raporu değildir: senaryonun bugünkü halini, yani hangi
parçalardan oluştuğunu, adım adım ne yaptığını, hangi parametrelerle sürüldüğünü
ve nasıl çalıştırıldığını anlatır. Gerekçeler ve bulgular ayrı belgededir
(insan_robot_isbirligi_raporu.docx).

Tur adımları senaryo kaynağından, parametreler launch dosyasından üretim anında
okunur.
"""
import ast
import os
import sys

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
PKG = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import hrc_figurler as F  # noqa: E402

OUT = os.path.join(HERE, "insan_robot_isbirligi_sistem_tanimi.docx")
LAUNCH = os.path.join(PKG, "launch", "human_robot_collaboration_scenario.launch.py")
FIG_WIDTH = Inches(5.9)
INK = "#22252a"


def h1(doc, t):
    doc.add_heading(t, level=1)


def h2(doc, t):
    doc.add_heading(t, level=2)


def p(doc, t):
    return doc.add_paragraph(t)


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def code(doc, text):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    par.paragraph_format.space_after = Pt(8)


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for c, name in zip(t.rows[0].cells, header):
        c.text = ""
        r = c.paragraphs[0].add_run(str(name))
        r.bold = True
        r.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for c, val in zip(cells, row):
            c.text = ""
            r = c.paragraphs[0].add_run(str(val))
            r.font.size = Pt(8.5)
    if widths:
        t.autofit = False
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph()


def figure(doc, filename, caption):
    path = os.path.join(HERE, filename)
    if not os.path.exists(path):
        raise FileNotFoundError(f"{path} yok — önce: python3 docs/hrc_figurler.py")
    doc.add_picture(path, width=FIG_WIDTH)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def launch_args(path):
    tree = ast.parse(open(path).read())
    out = []
    for node in ast.walk(tree):
        if not (isinstance(node, ast.Call)
                and getattr(node.func, "id", "") == "DeclareLaunchArgument"):
            continue
        name = node.args[0].value if node.args and isinstance(node.args[0], ast.Constant) \
            else "?"
        kw = {k.arg: k.value for k in node.keywords}

        def lit(key):
            if key not in kw:
                return ""
            try:
                return ast.literal_eval(kw[key])
            except Exception:
                return "<dinamik>"
        out.append((name, lit("default_value"), " ".join(str(lit("description")).split())))
    return out


def short(text, n=185):
    return text if len(text) <= n else text[:n - 1] + "…"


# --------------------------------------------------------------------------- #
def fig_architecture():
    boxes = [
        (0.3, 5.3, 3.4, 1.0, "human_robot_collaboration_scenario\n(tek düğüm: tur "
         "listesi + durum akışı)", "#fff4e6", "#d9480f"),
        (4.6, 6.1, 3.4, 0.8, "move_group (gripper'lı yapılandırma)\nplanlama + sahne",
         "#f3f0ff", "#7048e8"),
        (4.6, 4.9, 3.4, 0.8, "scaled_joint_trajectory_controller\nUR10e + Festo rayı",
         "#fff5f5", "#c92a2a"),
        (4.6, 3.7, 3.4, 0.8, "/gripper_controller\nOnRobot 2FG7 (MoveIt atlanır)",
         "#fff5f5", "#c92a2a"),
        (4.6, 2.5, 3.4, 0.8, "io_and_status_controller\nset_io · io_states (RTDE)",
         "#e7f5ff", "#1f6feb"),
        (0.3, 3.6, 3.4, 0.8, "Planning Scene\nvidalama aleti attach/detach",
         "#f8f9fa", "#868e96"),
        (0.3, 2.1, 3.4, 0.8, "Operatör: YEŞİL buton (DIN7)", "#ebfbee", "#2f9e44"),
        (0.3, 0.9, 3.4, 0.8, "Vidalama motoru (DOUT0)", "#fff9db", "#e8590c"),
    ]
    arrows = [((3.7, 6.0), (4.6, 6.5)), ((3.7, 5.8), (4.6, 5.3)),
              ((3.7, 5.5), (4.6, 4.1)), ((3.7, 5.35), (4.6, 3.1)),
              ((2.0, 5.3), (2.0, 4.4)),
              ((3.7, 2.5), (4.6, 2.8)),      # buton -> dijital giriş
              ((4.6, 2.6), (3.7, 1.3))]      # dijital çıkış -> motor
    fig, ax = plt.subplots(figsize=(8.4, 4.6), dpi=180)
    ax.set_xlim(0, 8.3)
    ax.set_ylim(0.5, 7.2)
    ax.axis("off")
    for x, y, w, hgt, label, fc, ec in boxes:
        ax.add_patch(FancyBboxPatch((x, y), w, hgt,
                                    boxstyle="round,pad=0.06,rounding_size=0.12",
                                    facecolor=fc, edgecolor=ec, lw=1.3))
        ax.text(x + w / 2, y + hgt / 2, label, ha="center", va="center", fontsize=7.8,
                color=INK)
    for a, b in arrows:
        ax.add_patch(FancyArrowPatch(a, b, arrowstyle="-|>", mutation_scale=11,
                                     color="#868e96", lw=1.1))
    ax.text(4.05, 6.62, "kol hedefleri", fontsize=6.8, color="#868e96", ha="center")
    ax.text(4.15, 4.32, "tırnak hedefi", fontsize=6.8, color="#868e96", ha="center")
    ax.text(4.15, 3.35, "set_io / io_states", fontsize=6.8, color="#868e96", ha="center")
    ax.text(4.15, 2.86, "DIN7", fontsize=6.8, color="#2f9e44", ha="center")
    ax.text(4.15, 1.62, "DOUT0", fontsize=6.8, color="#e8590c", ha="center")
    ax.set_title("İnsan-robot işbirlikli vidalama — düğüm, kontrolcüler ve operatör",
                 fontsize=10.5, color=INK)
    fig.tight_layout()
    out = os.path.join(HERE, "fig_hrc_mimari.png")
    fig.savefig(out, bbox_inches="tight")
    plt.close(fig)
    print("yazıldı:", out)


# --------------------------------------------------------------------------- #
def build():
    fig_architecture()
    cell = F.Cell(F.urdf_path())
    _, wps, seq = F.scenario_data()
    steps = F.seq_steps(seq)
    moves = [s for s in steps if s["kind"] == "move"]
    events = [s for s in steps if s["kind"] == "event"]
    n_button = sum(1 for e in events if e["key"] == "wait_green_button")
    args = launch_args(LAUNCH)
    P = np.array([F.tcp_of(cell, s["q"]) for s in moves])
    path_len = float(np.linalg.norm(np.diff(P, axis=0), axis=1).sum())

    doc = Document()
    doc.add_heading("İnsan-Robot İşbirlikli Vidalama — Sistem Tanımı", 0)
    sub = doc.add_paragraph("UR10e + OnRobot 2FG7 + operatör onaylı vidalama — "
                            "Senaryonun Bugünkü Hali")
    sub.style = doc.styles["Subtitle"]
    meta = doc.add_paragraph()
    meta.add_run("12 Eylül 2026    |    Cem Süha Yılmaz    |    "
                 "ROS 2 Humble / MoveIt 2").bold = True
    p(doc, "Bu belge senaryonun ŞU ANKİ halini anlatır: hangi parçalardan oluştuğu, "
           "turun adım adım ne yaptığı, hangi parametrelerle sürüldüğü, operatörle "
           "nasıl konuştuğu ve nasıl çalıştırıldığı. Karar geçmişi, ölçümler ve "
           "bulgular ayrı belgededir (insan_robot_isbirligi_raporu.docx). Tablolar "
           "senaryo kaynağından ve launch dosyasından üretim anında okunmuştur.")

    # 1
    h1(doc, "1. Sistem Tek Paragrafta")
    p(doc, "Robot bir montaj parçasının dört vidasını sıkar, ama her vida operatörün "
           "onayıyla başlar. Tur şöyle işler: robot vidalama aletini standından alır, "
           "besleyiciden vidayı alır, vidanın üstüne konumlanır ve DURUR. Operatör "
           "yeşil butona bastığında vidalama motoru dijital çıkıştan çalışır, robot "
           "on kat yavaş bir hızla sıkma derinliğine iner, motor kapatılır ve kol "
           "geri çekilir. Dört vida bittiğinde alet standına bırakılır ve kol home "
           "pozuna döner. Senaryo launch başına TEK tur koşar ve kendi kendine kapanır.")
    table(doc, ["Bileşen", "Bugünkü değer"], [
        ["Kol / grup", "UR10e, MoveIt grubu real_ur10e (Festo rayı + 6 eklem)"],
        ["Uç eleman", "OnRobot 2FG7, ur10e_gripper_joint (prizmatik)"],
        ["Ray konumu", "sabit 1.85 m (bu senaryoda ray hareket etmez)"],
        ["Tur uzunluğu", f"{len(steps)} adım ({len(moves)} hareket, {len(events)} olay)"],
        ["Vida sayısı", f"{n_button} (her biri ayrı operatör onayı ister)"],
        ["Uç yolu", f"{path_len:.2f} m"],
        ["Seyir / vidalama hız ölçeği", "0.1 / 0.01"],
        ["Planlayıcı", "RRTConnectkConfigDefault, 5 s, 10 deneme"],
    ], widths=[1.9, 4.0])
    figure(doc, "fig_hrc_cell.png",
           "Şekil 1: Hücrenin tamamı ve vidalama iş istasyonu. Vidalama aleti standı, "
           "vida besleyici, operatör mankeni ve ışık perdesi hücrenin URDF'inde ayrı "
           "linklerdir.")

    # 2
    h1(doc, "2. Mimari")
    figure(doc, "fig_hrc_mimari.png",
           "Şekil 2: Senaryo tek bir düğümdür; kol hareketleri MoveIt üzerinden, "
           "gripper doğrudan kendi kontrolcüsünden, operatör iletişimi ise UR "
           "kontrol kutusunun dijital giriş/çıkışları üzerinden yürür.")
    table(doc, ["Parça", "Sorumluluk"], [
        ["examples/human_robot_collaboration_scenario.py",
         "Senaryonun tamamı: waypoint listesi, tur akışı, gripper, planning scene "
         "attach/detach, GPIO ve operatör beklemesi."],
        ["launch/human_robot_collaboration_scenario.launch.py",
         "Parametreleri tipli biçimde geçirir ve düğüm bitince launch'ı kapatır "
         "(on_exit=Shutdown)."],
        ["pymoveit2_real (kütüphane)",
         "MoveIt 2 sarmalayıcı: poz/eklem hedefleri, 2π açma, en yakın dal IK'sı, "
         "en kısa yörünge seçimi."],
        ["move_group + real_ifarlab_gripper_moveit_config",
         "Planlama, çarpışma denetimi ve yürütme."],
        ["io_and_status_controller",
         "Dijital çıkış yazma (set_io) ve giriş okuma (io_states); RTDE üzerinden."],
    ], widths=[2.4, 3.5])

    # 3
    h1(doc, "3. Turun Adımları")
    p(doc, "Tur, tek bir listeden yürür. Liste hem eklem hedeflerini hem de sözlük "
           "biçiminde komutları taşır; aşağıdaki tablo o listenin bugünkü halidir.")
    rows = []
    for i, s in enumerate(steps, 1):
        if s["kind"] == "move":
            name = F.name_of(wps, s["q"])
            kind = "hareket (yavaş)" if s["speed"] else "hareket"
            detail = f"hız ölçeği {s['speed']}" if s["speed"] else "seyir hızı"
            rows.append([i, kind, name, detail])
        else:
            key, val = s["key"], s["value"]
            label = {"gripper_position": "gripper",
                     "attach_screwdriver": "planning scene",
                     "detach_screwdriver": "planning scene",
                     "screwdriver": "vidalama çıkışı",
                     "wait_green_button": "operatör",
                     "wait": "bekleme"}.get(key, key)
            detail = {"gripper_position": f"{val} m",
                      "attach_screwdriver": "aleti attach et",
                      "detach_screwdriver": "aleti detach et",
                      "screwdriver": "DOUT0 " + ("ON" if val else "OFF"),
                      "wait_green_button": "yeşil butona basılmasını bekle",
                      "wait": f"{val} s"}.get(key, str(val))
            rows.append([i, label, "-", detail])
    table(doc, ["#", "Tür", "Waypoint", "Ayrıntı"], rows,
          widths=[0.3, 1.2, 1.3, 3.1])
    figure(doc, "fig_hrc_postures.png",
           "Şekil 3: Turun dört anahtar pozu (URDF'ten ileri kinematik).")
    figure(doc, "fig_hrc_timeline.png",
           "Şekil 4: Turun zaman çizgisi. Hareket süreleri eklem limitlerinden "
           "kestirilmiştir; operatör beklemesi varsayımdır.")
    p(doc, "Akışın üç kuralı vardır: (1) her tur, gripper'ı TAM AÇIK konuma "
           "getirerek başlar, çünkü önceki tur yarıda kesilmiş olabilir; (2) alet "
           "planning scene'e gripper KAPANDIKTAN sonra eklenir ve gripper "
           "AÇILMADAN önce çıkarılır; (3) her vidada motor, kol geri çıkmadan ÖNCE "
           "kapatılır.")

    # 4
    h1(doc, "4. Operatörle İletişim (GPIO)")
    figure(doc, "fig_hrc_gpio.png",
           "Şekil 5: Pin haritası ve bir vidanın el sıkışma sırası.")
    table(doc, ["Pin", "İşlev", "Senaryodaki rolü"], [
        ["standard_digital_out[0]", "vidalama SIKMA", "kullanılıyor"],
        ["standard_digital_out[1]", "vidalama SÖKME", "rezerve (reverse=True)"],
        ["standard_digital_in[7]", "YEŞİL buton", "kullanılıyor"],
        ["standard_digital_in[6]", "KIRMIZI buton", "rezerve"],
        ["standard_digital_in[5]", "BEYAZ buton", "rezerve"],
    ], widths=[1.9, 1.6, 2.4])
    bullets(doc, [
        "Buton YÜKSELEN KENARLA okunur; çağrı anında basılıysa önce bırakılması "
        "beklenir, böylece tek basış iki vidalamayı tetiklemez.",
        "green_button_timeout varsayılanı 0.0'dır: robot süresiz bekler, kendi "
        "başına vidalamaz.",
        "gpio_mode=auto iken düğüm robot_description'a bakar; mock/sim donanım "
        "görürse GPIO'yu kapatır, vidalama adımlarını atlar ve buton beklemez. "
        "force_on / force_off bu kararı geçersiz kılar.",
        "io_states akışı durursa girişler sessizce donar; düğüm bunu 'bayat' diye "
        "raporlar ve pin hiç yayınlanmıyorsa ayrı bir hata basar.",
    ])

    # 5
    h1(doc, "5. Yapılandırma Referansı")
    table(doc, ["Argüman", "Varsayılan", "Açıklama"],
          [[a, str(d), short(desc)] for a, d, desc in args], widths=[1.5, 0.9, 3.5])
    p(doc, "Launch'ta görünmeyen ama düğümde tanımlı olan diğer ayarlar: gripper "
           "konumları (tam açık 0.0, bırakma 0.003, kapalı 0.026 m), gripper oturma "
           "süresi (0.7 s), vidalama sonrası bekleme (adım listesinde 1.5 s), "
           "adımlar arası sabit bekleme (0.5 s) ve nokta başına yeniden deneme "
           "sayısı (3).")

    # 6
    h1(doc, "6. Çalıştırma")
    code(doc,
         "# 1) hücre + sürücü + MoveIt (gripper'lı yapılandırma)\n"
         "ros2 launch my_robot_cell_control hil_test_whole_unified.launch.py \\\n"
         "    use_gripper:=true use_fake_hardware:=false\n\n"
         "# 2) senaryo (tek tur koşar, biter ve kapanır)\n"
         "ros2 launch pymoveit2_real human_robot_collaboration_scenario.launch.py\n\n"
         "# simülasyonda denemek (GPIO kendiliğinden kapanır)\n"
         "ros2 launch pymoveit2_real human_robot_collaboration_scenario.launch.py \\\n"
         "    fake_button_delay:=2.0")
    p(doc, "Arayüzdeki 'Human-Robot Collaboration Scenario' düğmesi HIL'i "
           "use_gripper:=true ile kaldırır ve senaryoyu başlatır. Bu senaryo kendi "
           "launch'ında on_exit=Shutdown taşıdığı için, diğer senaryoların aksine, "
           "bittiğinde süreç normal biçimde sonlanır.")

    # 7
    h1(doc, "7. Çalışma Kuralları ve Sınırlar")
    bullets(doc, [
        "Vidalama açık döngüdür: sabit derinliğe inilir, 1.5 s beklenir ve motor "
        "kapatılır. Tork veya derinlik geri beslemesi okunmaz.",
        "Aleti kavrama pozunda gripper ile tezgâh arasında yalnızca birkaç "
        "milimetre pay vardır; kol linklerine büyük padding verilirse bu poz "
        "planlanamaz hale gelir.",
        "Planning scene'e eklenen alet temsili küçük bir silindirdir; alet "
        "taşınırken serbest planlamaya izin verilecekse büyütülmelidir.",
        "Operatör mankeni URDF'te vardır ama SRDF onu kolun bütün hareketli "
        "linkleriyle çarpışma denetiminden muaf tutar. Güvenlik fiziksel katmandan "
        "gelir: ışık perdesi, bariyerler ve robotun onay beklemesi.",
        "Senaryo launch başına tek tur koşar; yeni tur için launch yeniden "
        "başlatılır.",
        "Kapanış yolu koşulsuzdur: istisna ya da Ctrl+C dahil her çıkışta önce "
        "vidalama çıkışları düşürülür, sonra kol home pozuna gider.",
    ])

    doc.save(OUT)
    print("yazıldı:", OUT)


if __name__ == "__main__":
    build()
