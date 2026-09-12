#!/usr/bin/env python3
"""multirobot_viewpoint_planner paketinin BUGÜNKÜ halini anlatan sistem tanımını üretir.

    python3 docs/sistem_tanimi_uret.py

Bu belge GELİŞME raporu değildir: neyin ne zaman değiştiğini değil, sistemin şu
anda nasıl çalıştığını anlatır. Karar/gelişme kaydı ayrı dosyadadır
(viewpoint_inspection_system_report.docx, rapor_uret.py ile üretilir).

Parametre tabloları launch dosyasından, veri şemaları güncel plan dosyasından,
başarım sayıları octomap çıktılarından üretim anında okunur.
"""
import ast
import json
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
PKG = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import figur_sasi as FS  # noqa: E402

OUT = os.path.join(HERE, "multirobot_sistem_tanimi.docx")
LAUNCH = os.path.join(PKG, "launch", "multirobot_inspection.launch.py")
PLAN = os.path.join(PKG, "plans", "multirobot_viewpoint_plan.json")
FIG_WIDTH = Inches(5.9)
INK = "#22252a"


def h1(doc, t):
    doc.add_heading(t, level=1)


def h2(doc, t):
    doc.add_heading(t, level=2)


def p(doc, t):
    return doc.add_paragraph(t)


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def code(doc, text):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    par.paragraph_format.space_after = Pt(8)


def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for c, name in zip(t.rows[0].cells, header):
        c.text = ""
        r = c.paragraphs[0].add_run(str(name))
        r.bold = True
        r.font.size = Pt(9)
    for row in rows:
        cells = t.add_row().cells
        for c, val in zip(cells, row):
            c.text = ""
            r = c.paragraphs[0].add_run(str(val))
            r.font.size = Pt(8.5)
    if widths:
        t.autofit = False
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    doc.add_paragraph()


def figure(doc, filename, caption):
    path = os.path.join(HERE, filename)
    if not os.path.exists(path):
        raise FileNotFoundError(f"{path} yok — önce figür üreticilerini çalıştırın")
    doc.add_picture(path, width=FIG_WIDTH)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def launch_args(path):
    tree = ast.parse(open(path).read())
    out = []
    for node in ast.walk(tree):
        if not (isinstance(node, ast.Call)
                and getattr(node.func, "id", "") == "DeclareLaunchArgument"):
            continue
        name = node.args[0].value if node.args and isinstance(node.args[0], ast.Constant) \
            else "?"
        kw = {k.arg: k.value for k in node.keywords}

        def lit(key):
            if key not in kw:
                return ""
            try:
                return ast.literal_eval(kw[key])
            except Exception:
                return "<dinamik>"
        out.append((name, lit("default_value"), " ".join(str(lit("description")).split())))
    return out


def short(text, n=185):
    return text if len(text) <= n else text[:n - 1] + "…"


# --------------------------------------------------------------------------- #
def fig_architecture():
    boxes = [
        (0.3, 6.4, 3.2, 0.9, "Şasi mesh'i (STL)", "#e7f5ff", "#1f6feb"),
        (0.3, 5.0, 3.2, 1.0, "multirobot_planner_node\nortak aday havuzu + tahsis",
         "#fff4e6", "#d9480f"),
        (0.3, 3.6, 3.2, 1.0, "plans/multirobot_viewpoint_plan.json\nur_viewpoints + "
         "kawasaki_viewpoints", "#f8f9fa", "#868e96"),
        (0.3, 1.9, 1.5, 1.1, "ur_inspection\n_node", "#ebfbee", "#2f9e44"),
        (2.0, 1.9, 1.5, 1.1, "kawasaki_\ninspection_node", "#ebfbee", "#2f9e44"),
        (0.3, 0.4, 3.2, 0.9, "pcds/{real,sim}_pcds/{ur,kawasaki}_data → octomap",
         "#f8f9fa", "#868e96"),
        (4.4, 5.0, 3.6, 1.0, "move_group (whole_cell_hw)\nIK · plan · sahne",
         "#f3f0ff", "#7048e8"),
        (4.4, 3.3, 3.6, 0.9, "scaled_joint_trajectory_controller\n(UR10e + ray)",
         "#fff5f5", "#c92a2a"),
        (4.4, 1.9, 3.6, 0.9, "/kawasaki/kawasaki_controller\n(Kawasaki + world_to_agv)",
         "#fff5f5", "#c92a2a"),
        (4.4, 0.4, 3.6, 0.9, "agv bridge → rosbridge → ROS 1\n(AGV, sabit 0.05 m/s)",
         "#fff9db", "#e8590c"),
    ]
    arrows = [((1.9, 6.4), (1.9, 6.0)), ((1.9, 5.0), (1.9, 4.6)),
              ((1.9, 3.6), (1.9, 3.0)), ((1.05, 1.9), (1.05, 1.3)),
              ((2.75, 1.9), (2.75, 1.3)),
              ((3.5, 5.5), (4.4, 5.5)), ((4.4, 5.2), (3.5, 5.2)),
              ((1.82, 2.55), (4.4, 3.75)), ((3.5, 2.2), (4.4, 2.35)),
              ((6.2, 1.9), (6.2, 1.3))]
    fig, ax = plt.subplots(figsize=(8.4, 4.8), dpi=180)
    ax.set_xlim(0, 8.3)
    ax.set_ylim(0, 7.7)
    ax.axis("off")
    for x, y, w, hgt, label, fc, ec in boxes:
        ax.add_patch(FancyBboxPatch((x, y), w, hgt,
                                    boxstyle="round,pad=0.06,rounding_size=0.12",
                                    facecolor=fc, edgecolor=ec, lw=1.3))
        ax.text(x + w / 2, y + hgt / 2, label, ha="center", va="center", fontsize=7.8,
                color=INK)
    for a, b in arrows:
        ax.add_patch(FancyArrowPatch(a, b, arrowstyle="-|>", mutation_scale=11,
                                     color="#868e96", lw=1.1))
    ax.text(3.95, 5.75, "IK + çarpışma", fontsize=7, color="#868e96", ha="center")
    ax.text(3.05, 3.25, "yörünge (UR)", fontsize=7, color="#868e96", ha="left")
    ax.text(3.6, 2.5, "yörünge (Kawasaki)", fontsize=7, color="#868e96", ha="left")
    ax.text(6.55, 1.52, "world_to_agv", fontsize=7, color="#868e96", ha="left")
    ax.text(1.9, 1.55, "iki kol AYRI SÜREÇ, eşzamanlı koşar", fontsize=7,
            color="#2f9e44", ha="center")
    ax.set_title("multirobot_viewpoint_planner — veri akışı ve iki kolun ayrımı",
                 fontsize=10.5, color=INK)
    fig.tight_layout()
    out = os.path.join(HERE, "fig_multirobot_mimari.png")
    fig.savefig(out, bbox_inches="tight")
    plt.close(fig)
    print("yazıldı:", out)


# --------------------------------------------------------------------------- #
def build():
    fig_architecture()
    plan = json.load(open(PLAN))
    ur, kawa = plan["ur_viewpoints"], plan["kawasaki_viewpoints"]
    args = launch_args(LAUNCH)
    P = FS.PCDS
    cov_real = FS.coverage(f"{P}/real_pcds/beliefMap_real.ot",
                           f"{P}/real_pcds/occupancyMap_real.ot")
    cov_sim = FS.coverage(f"{P}/sim_pcds/beliefMap_sim.ot",
                          f"{P}/sim_pcds/occupancyMap_sim.ot")
    single = FS.coverage(
        f"{P}/single_ur10e/real_data/beliefMap_single_ur10e_real.ot",
        f"{P}/single_ur10e/real_data/occupancyMap_single_ur10e_real.ot")

    doc = Document()
    doc.add_heading("multirobot_viewpoint_planner — Sistem Tanımı", 0)
    sub = doc.add_paragraph("İşbirlikli UR10e + Kawasaki RS005L ile Şasi Muayenesi — "
                            "Paketin Bugünkü Hali")
    sub.style = doc.styles["Subtitle"]
    meta = doc.add_paragraph()
    meta.add_run("12 Eylül 2026    |    Cem Süha Yılmaz    |    "
                 "ROS 2 Humble / MoveIt 2").bold = True
    p(doc, "Bu belge paketin ŞU ANKİ halini anlatır: parçaları, uçtan uca akışı, "
           "arayüzleri, yapılandırması, ürettiği dosyalar ve bugünkü ölçülen "
           "başarımı. Değişiklik gerekçeleri ve karar geçmişi ayrı belgededir "
           "(viewpoint_inspection_system_report.docx).")

    # 1
    h1(doc, "1. Sistem Tek Paragrafta")
    p(doc, "Sistem, tek bir şasiyi İKİ KOLLA muayene eder. Aynı aday havuzundan "
           "üretilen bakış-noktaları kollara paylaştırılır: her bakış-noktası, ona "
           "ulaşabilen kola verilir. Böylece tek kolun asla göremeyeceği yüzler de "
           "(rayın öbür tarafında kalanlar) kaplamaya girer ve iki kol eşzamanlı "
           "çalıştığı için tur duvar saatiyle kısalır. UR10e 2 metrelik Festo rayı "
           "üzerinde, Kawasaki ise AGV rayı üzerinde hareket eder; ikisi de kendi "
           "sürecinde, kendi kayıtlı yörüngeleriyle koşar.")
    table(doc, ["Bileşen", "Bugünkü değer"], [
        ["Kollar", "UR10e (real_ur10e) + Kawasaki RS005L (real_kawasaki)"],
        ["Doğrusal eksenler", "Festo rayı (ur10e_base_to_robot_mount), AGV (world_to_agv)"],
        ["Sensörler", "UR: /sick_points · Kawasaki: /kawasaki/pointcloud "
                      "(sim: /sim/pointcloud, /sim/kawasaki/pointcloud)"],
        ["Plandaki bakış-noktası", f"{len(ur)} UR + {len(kawa)} Kawasaki = "
                                   f"{len(ur) + len(kawa)}"],
        ["Planlayıcı tahmini kaplama", f"%{100 * plan['coverage_achieved']:.1f}"],
        ["Gerçek robot octomap kaplaması", f"%{100 * cov_real['frac']:.1f}"],
        ["Simülasyon kaplaması", f"%{100 * cov_sim['frac']:.1f}"],
        ["Tek kola göre kazanç",
         f"+{100 * (cov_real['frac'] - single['frac']):.1f} puan "
         f"(tek kol %{100 * single['frac']:.1f})"],
    ], widths=[1.9, 4.0])

    # 2
    h1(doc, "2. Mimari")
    figure(doc, "fig_multirobot_mimari.png",
           "Şekil 1: Paketin veri akışı. Planlama tek düğümde, yürütme iki AYRI "
           "süreçte yapılır; AGV, Kawasaki kontrolcüsünün arkasındaki ayrı ve yavaş "
           "bir platformdur.")
    table(doc, ["Modül", "Sorumluluk"], [
        ["multirobot_planner_node.py",
         "Hedef örnekleme, aday üretimi, görünürlük, kol tahsisi, sıralama ve plan "
         "yazımı. Yürütücünün sahnesini (zemin + padding) IK'dan önce kurar."],
        ["robot_allocator.py",
         "Ortak aday havuzunu greedy biçimde kollara paylaştırır. İki hedefi vardır: "
         "kaplamayı tek kolun tavanının üstüne çıkarmak ve turu kısaltmak. Kol başına "
         "kap ve dengeleme anahtarları buradadır."],
        ["inspection_base.py",
         "İki kol için ortak yürütme mantığı: plan okuma, trajectory cache, 2π açma, "
         "padding, sahne kurulumu, ölçülen-varış denetimi, bulut yakalama."],
        ["ur_inspection_node.py",
         "UR'ye özgü katman: pose-goal seçeneği, ur10e_ padding, UR kamera konuları."],
        ["kawasaki_inspection_node.py",
         "Kawasaki'ye özgü katman: yörüngeyi doğrudan /kawasaki/kawasaki_controller'a "
         "gönderir, ölçülen-varış kapısını kullanır, AGV'nin yavaşlığını tolere eder."],
        ["visualization.py / plan_visualizer.py",
         "RViz işaretçileri ve plan görselleştirmesi."],
    ], widths=[1.7, 4.2])
    p(doc, "Düğüm ayrımının sebebi basittir: iki kol eşzamanlı hareket etmelidir ve "
           "biri diğerinin planlama çağrısını beklememelidir. Ortak mantık tek bir "
           "temel sınıfta durur, kola özgü olan ince alt sınıflarda kalır.")

    # 3
    h1(doc, "3. Uçtan Uca Akış")
    table(doc, ["#", "Adım", "Çıktısı"], [
        ["1", "Mesh örneklenir, adaylar üretilir, görünürlük hesaplanır",
         "ortak aday havuzu"],
        ["2", "Yürütücünün sahnesi kurulur (zemin + padding)",
         "IK, yürütmeyle aynı dünyada sınanır"],
        ["3", "Her aday için iki kolun da IK'sı denenir",
         "kol başına ulaşılabilir küme"],
        ["4", "Greedy tahsis: bakış-noktası, ona ulaşan kola verilir",
         "ur_viewpoints + kawasaki_viewpoints"],
        ["5", "Her kolun turu Y bantlarına göre sıralanır", "gezme sırası"],
        ["6", "Plan JSON'a yazılır (ik_scene damgasıyla)",
         "plans/multirobot_viewpoint_plan.json"],
        ["7", "İki yürütücü süreci başlar, her biri başlangıç pozuna gider", "-"],
        ["8", "Her durak: cache'ten oynat ya da planla ve kaydet",
         "plans/trajectories/{ur,kawasaki}_<id>.json"],
        ["9", "Durakta bulut + poz kaydedilir",
         "pcds/{real,sim}_pcds/{ur,kawasaki}_data"],
        ["10", "Turlar bitince kollar home pozuna döner", "-"],
        ["11", "Bulutlar octomap'e çevrilir", "occupancyMap_*.ot / beliefMap_*.ot"],
    ], widths=[0.3, 3.0, 2.6])
    figure(doc, "fig_multirobot_plan.png",
           f"Şekil 2: Güncel plan — {len(ur)} UR + {len(kawa)} Kawasaki "
           "bakış-noktası. Tepe görünümünde iki kümenin X aralıklarının hiç "
           "örtüşmemesi, işbirliğinin uzaysal karşılığıdır.")

    # 4
    h1(doc, "4. Arayüzler")
    table(doc, ["Yön", "Ad", "Not"], [
        ["abone", "/sick_points", "UR'nin gerçek sensörü"],
        ["abone", "/kawasaki/pointcloud", "Kawasaki'nin gerçek sensörü"],
        ["abone", "/sim/pointcloud, /sim/kawasaki/pointcloud", "Gazebo karşılıkları"],
        ["abone", "/joint_states", "iki kolun ölçülen durumu (varış denetimi)"],
        ["servis", "/compute_ik, /plan_kinematic_path", "tahsis ve yürütme planlaması"],
        ["servis", "/apply_planning_scene, /get_planning_scene", "zemin + padding"],
        ["action", "/scaled_joint_trajectory_controller/follow_joint_trajectory",
         "UR yörüngesi"],
        ["action", "/kawasaki/kawasaki_controller/follow_joint_trajectory",
         "Kawasaki + AGV yörüngesi (world_to_agv dahil)"],
        ["dolaylı", "/agv/goal_position → rosbridge → ROS 1 platformu",
         "AGV sabit 0.05 m/s ile ve kendi eylem sunucusuyla hareket eder"],
    ], widths=[0.7, 2.7, 2.5])
    p(doc, "AGV'nin ayrı ve yavaş bir platform olması, bu paketin en çok kural "
           "gerektiren yeridir: Kawasaki kontrolcüsünün goal_time'ı büyük tutulur "
           "(240 s), yürütücü ölçülen-varışla bekler (arrival_timeout_sec 180 s) ve "
           "move_group üzerinden gönderilen hareketler için kontrolcüye özel süre "
           "marjı tanımlıdır. Muayene yürütücüsü yörüngeyi doğrudan kontrolcüye "
           "gönderdiği için move_group'un süre denetimine takılmaz.")

    # 5
    h1(doc, "5. Yapılandırma Referansı")
    p(doc, f"Yürütme launch dosyası {len(args)} argüman tanımlar:")
    table(doc, ["Argüman", "Varsayılan", "Açıklama"],
          [[a, str(d), short(desc)] for a, d, desc in args], widths=[1.5, 0.9, 3.5])
    p(doc, "Planlama tarafı config/multirobot_params.yaml ile sürülür; en çok "
           "dokunulan anahtarlar:")
    table(doc, ["Anahtar", "Değer", "Ne yapar"], [
        ["min_marginal_coverage", "0.0035", "greedy durma ölçütü"],
        ["max_incidence_angle_deg", "85", "görülmüş sayılma açısı"],
        ["max_viewpoints_ur / _kawasaki", "45 / 13",
         "kol başına kap; Kawasaki bilerek dar, çünkü her durağı AGV'yi de hareket "
         "ettirir"],
        ["order_mode / order_band_width_m", "y_bands / 0.30",
         "her kol şasiyi Y bantları hâlinde tarar"],
        ["collision_padding", "0.04 m", "ur10e_* ve link1..6 gövdelerine pay"],
        ["ik_scene_setup", "true", "planlayıcı, yürütücünün sahnesini kurar"],
    ], widths=[1.9, 0.9, 3.1])

    # 6
    h1(doc, "6. Üretilen Dosyalar")
    vp_keys = ", ".join(ur[0].keys())
    table(doc, ["Dosya", "İçerik"], [
        ["plans/multirobot_viewpoint_plan.json",
         f"üst düzey: {', '.join(plan.keys())}"],
        ["… ur_viewpoints[] / kawasaki_viewpoints[]", f"her kayıt: {vp_keys}"],
        ["… ik_scene", "planın doğrulandığı sahne (zemin, padding, link sayısı)"],
        ["… manual_edits", "plan üzerinde elle yapılan düzeltmelerin kaydı"],
        ["plans/trajectories/ur_<id>.json, kawasaki_<id>.json",
         "kayıt-oynat yörüngeleri (dosya adı kol etiketi + viewpoint kimliği)"],
        ["pcds/{real,sim}_pcds/{ur,kawasaki}_data/{pcds,poses}",
         "kol başına bulut ve poz kayıtları"],
        ["occupancyMap_*.ot / beliefMap_*.ot", "ölçülen ve referans voksel haritaları"],
    ], widths=[2.4, 3.5])

    # 7
    h1(doc, "7. Çalıştırma")
    code(doc,
         "# plan üretimi\n"
         "ros2 launch multirobot_viewpoint_planner multirobot_planning.launch.py\n\n"
         "# iki kollu muayene turu (gerçek robotlar)\n"
         "ros2 launch multirobot_viewpoint_planner multirobot_inspection.launch.py \\\n"
         "    only_sim:=false\n\n"
         "# tek kolu koşturmak\n"
         "ros2 launch multirobot_viewpoint_planner multirobot_inspection.launch.py \\\n"
         "    only_ur:=true        # ya da only_kawasaki:=true\n\n"
         "# kayıtlı yörüngeleri yeniden üret\n"
         "ros2 launch multirobot_viewpoint_planner multirobot_inspection.launch.py \\\n"
         "    force_replan:=true")
    p(doc, "Arayüzdeki 'Multi-Robot Inspection Scenario' düğmesi aynı zinciri kurar. "
           "Launch süreci turun bitmesiyle sonlanmaz (kalıcı görselleştirici düğümü "
           "yüzünden); bitişin ölçütü iki yürütücü sürecinin sonlanmasıdır.")

    # 8
    h1(doc, "8. Bugünkü Ölçülen Başarım")
    table(doc, ["Koşu", "Şasi vokseli", "Kaplanan", "Kaplama"], [
        ["İki kol — gerçek robot", len(cov_real["belief"]), len(cov_real["covered"]),
         f"%{100 * cov_real['frac']:.1f}"],
        ["İki kol — simülasyon", len(cov_sim["belief"]), len(cov_sim["covered"]),
         f"%{100 * cov_sim['frac']:.1f}"],
        ["Tek kol — gerçek robot (karşılaştırma)", len(single["belief"]),
         len(single["covered"]), f"%{100 * single['frac']:.1f}"],
    ], widths=[2.4, 1.2, 1.1, 1.0])
    figure(doc, "fig_octomap_real_vs_sim.png",
           "Şekil 3: Kaplanan (yeşil) ve kaplanamayan (kırmızı) şasi vokselleri; "
           "gerçek ve simülasyon aynı kamera açılarıyla.")
    figure(doc, "fig_part_coverage.png",
           "Şekil 4: Parça başına kaplama. Gerçek robottaki eksik iki bölgede "
           "toplanıyor: en üst raylar ve AGV güverte seviyesindeki alt raylar.")

    # 9
    h1(doc, "9. Çalışma Kuralları ve Sınırlar")
    bullets(doc, [
        "Trajectory cache URDF/SRDF/padding değişikliklerine kördür; hücre "
        "geometrisi değiştiyse force_replan gerekir.",
        "Kayıtlı yörünge, kaydedildiği HIZLA oynatılır. Kawasaki hız/ivme ölçeğini "
        "değiştirmek eski kayıtları etkilemez.",
        "AGV zincirindeki üç zaman aşımı (kontrolcü goal_time, yürütücü "
        "arrival_timeout_sec, move_group süre denetimi) birlikte büyütülmelidir; "
        "yalnız birini büyütmek hareketi başka bir katmanda kestirir.",
        "İki kol aynı sahneyi paylaşır ama ayrı süreçlerde koşar; aralarında "
        "çarpışma önleme YOKTUR. Planlar uzaysal olarak ayrık olduğu için (UR ve "
        "Kawasaki şasinin zıt yüzlerinde) bugüne kadar gerekmedi.",
        "Kawasaki tarafında bakış-noktası sayısı bilerek düşük tutulur: her durak "
        "AGV'yi de hareket ettirir ve AGV 0.05 m/s ile gider.",
        "Gerçek robottaki kaplama eksiği rastgele değil, kolun erişim sınırındaki "
        "iki bölgede toplanır.",
    ])

    doc.save(OUT)
    print("yazıldı:", OUT)


if __name__ == "__main__":
    build()
